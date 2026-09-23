"""Build the presentable ATC normal-flow validation report."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "generated" / "ATC_ABM_Normal_Flow_Validation_Report.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(int(width_inches * 1440)))
    width.set(qn("w:type"), "dxa")
    tc_pr.append(width)


def style_table(table, widths, numeric_from=1):
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            set_width(cell, widths[c])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r == 0:
                set_cell_shading(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c >= numeric_from else WD_ALIGN_PARAGRAPH.LEFT
            else:
                if r % 2 == 0:
                    set_cell_shading(cell, PALE)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c >= numeric_from else WD_ALIGN_PARAGRAPH.LEFT
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(10)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for level, size in ((1, 15), (2, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("ATC Crowd Safety Explorer\nNormal Flow Validation Report")
    title_run.font.name = "Aptos Display"
    title_run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_run.font.size = Pt(25)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Observed trajectory forecasting with history based route intent and Social Force simulation")
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(68, 68, 68)
    subtitle.paragraph_format.space_after = Pt(18)

    add_body(doc, "This report validates whether the Agent Based Model can reproduce short-term pedestrian movement observed in the ATC shopping-center tracking data. The final held-out day shows a clear improvement when the model is seeded with a causal five-second trajectory history and uses that history to infer route intent. At a 10-second forecast horizon, mean final displacement error is 3.505 m, 29.71 percent lower than the constant-velocity baseline.")

    add_heading(doc, "Executive summary")
    summary = doc.add_table(rows=1, cols=3)
    headers = ["Question", "Answer", "Evidence"]
    for cell, value in zip(summary.rows[0].cells, headers):
        cell.text = value
    rows = [
        ("What was validated", "Normal-flow, short-term trajectory forecasting", "Six independent snapshots on a held-out ATC day"),
        ("What was compared", "History-aware Social Force forecast versus constant velocity", "All models start from the same observed people and positions"),
        ("Main outcome", "History improves every tested horizon", "17.74 percent to 33.00 percent FDE improvement versus baseline"),
        ("What this does not prove", "Real emergency or stampede prevention performance", "No ATC stampede event labels are available in this validation"),
    ]
    for values in rows:
        cells = summary.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(summary, [1.45, 2.15, 2.9], numeric_from=5)

    add_heading(doc, "System under test")
    add_body(doc, "The system begins with ATC person tracks at a user-selected date and time. Each track contains position, speed, motion angle, and facing angle. The model uses the observed people as ABM agents, constrains motion to the empirically walkable ATC footprint, and advances them with a calibrated Social Force model.")
    add_bullets(doc, [
        "Seed state: observed pedestrian position and identifier at the selected instant.",
        "Trajectory history: the five seconds immediately before that instant, used only as past evidence.",
        "Route intent: likely destination sampled from routes learned from other ATC days, weighted by smoothed heading and soft origin context.",
        "Movement model: map-constrained Social Force dynamics. The normal-flow calibration selected near-zero repulsion because it best matched this low-density dataset; facing-aware interaction remains enabled for dense safety counterfactuals.",
    ])

    add_heading(doc, "Validation design")
    add_body(doc, "The evaluation deliberately separates model selection from final scoring. Parameter selection was performed on 11 November 2012. Final performance was then measured on the unseen 14 November 2012 file. When evaluating either day, the route and speed profile is built from all available daily files except that evaluation day. This prevents information from the test day being used to construct its calibration profile.")
    design = doc.add_table(rows=1, cols=2)
    for cell, value in zip(design.rows[0].cells, ["Stage", "Procedure"]):
        cell.text = value
    for values in [
        ("1. Daily profile exclusion", "For a target day, aggregate speed and endpoint-route distributions only from the other nine available days."),
        ("2. Development calibration", "On 11 Nov, test 9 Social Force parameter combinations at 10 seconds across 3 snapshots. Select the lowest combined displacement and density objective."),
        ("3. Frozen configuration", "Selected values: relaxation time 10.0 s, repulsion strength 0.0, repulsion range 0.12 m, facing-aware enabled."),
        ("4. Held-out forecast", "On 14 Nov, choose 6 times spread through the day. Start from observed people at time t and forecast 1, 5, or 10 seconds."),
        ("5. Scoring", "Match simulated and actual people by persistent ATC person ID at t plus horizon. Score final displacement, speed, and 2 m grid density."),
    ]:
        cells = design.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(design, [1.7, 4.8], numeric_from=3)

    add_heading(doc, "Causal history and route intent inference")
    add_body(doc, "The improved configuration does not look ahead to the ground-truth future. For every tracked person, it fits a straight-line velocity estimate to the observations in the previous five seconds. The resulting heading and speed seed the agent. Candidate routes are taken from the held-out-day-excluded calibration distribution. A route receives greater weight when its destination aligns with the smoothed heading and when its learned origin is compatible with the recent path. The later ATC position is read only after simulation, for scoring.")

    add_heading(doc, "Metrics")
    metrics = doc.add_table(rows=1, cols=3)
    for cell, value in zip(metrics.rows[0].cells, ["Metric", "Definition", "Why it matters"]):
        cell.text = value
    for values in [
        ("Final displacement error FDE", "Mean Euclidean distance in metres between simulated and actual location at the horizon.", "Primary individual trajectory accuracy measure."),
        ("Speed MAE", "Mean absolute difference between simulated and actual speed in m/s at the horizon.", "Checks whether motion magnitude is plausible."),
        ("Density L1 error", "Difference in pedestrian counts across 2 m spatial cells, normalized by observed people.", "Tests whether the model reproduces local crowd concentration, not only individual endpoints."),
    ]:
        cells = metrics.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(metrics, [1.45, 3.25, 1.8], numeric_from=3)

    add_heading(doc, "Held out results")
    add_body(doc, "The table compares the same calibrated Social Force model without trajectory history against the completed history-aware configuration. The constant-velocity benchmark uses the instantaneous observed speed and direction and has no route, map, or interaction model. Lower errors are better.")
    results = doc.add_table(rows=1, cols=6)
    for cell, value in zip(results.rows[0].cells, ["Horizon", "No-history FDE", "History FDE", "Constant velocity FDE", "History gain vs baseline", "History density L1"]):
        cell.text = value
    for values in [
        ("1 second", "0.448 m", "0.377 m", "0.458 m", "17.74%", "0.473"),
        ("5 seconds", "2.024 m", "1.567 m", "2.338 m", "33.00%", "1.032"),
        ("10 seconds", "4.749 m", "3.505 m", "4.986 m", "29.71%", "1.364"),
    ]:
        cells = results.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(results, [0.8, 1.0, 0.9, 1.2, 1.2, 1.0])

    density = doc.add_table(rows=1, cols=4)
    for cell, value in zip(density.rows[0].cells, ["Horizon", "History density L1", "Constant velocity density L1", "Density error reduction"]):
        cell.text = value
    for values in [
        ("1 second", "0.473", "0.529", "10.59%"),
        ("5 seconds", "1.032", "1.441", "28.38%"),
        ("10 seconds", "1.364", "1.805", "24.43%"),
    ]:
        cells = density.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(density, [1.0, 1.7, 2.0, 1.8])

    add_body(doc, "Interpretation: the five-second trajectory history helps at all three horizons. The largest relative final-position gain occurs at five seconds, while the 10-second forecast retains a 29.71 percent improvement over the constant-velocity benchmark. The history-aware model also improves spatial density reproduction at every horizon. This is evidence that recent movement provides useful route-intent information beyond a single instantaneous heading.")

    add_heading(doc, "Calibration result")
    add_body(doc, "The selected normal-flow configuration was chosen only on the 11 Nov development day. It produced a development Social Force FDE of 3.020 m compared with 4.845 m for constant velocity, and density L1 of 1.362 compared with 1.790. The held-out result is therefore a separate estimate of generalization rather than a reused calibration score.")

    add_heading(doc, "Interpretation for crowd safety use")
    add_body(doc, "This validation supports the normal-operation component of the project: the system can initialize agents from live-like tracking data, infer short-term pedestrian intent from recent tracks, and forecast a more realistic near-future distribution than an unstructured straight-line baseline. That capability is a prerequisite for safety monitoring, because risk indicators must be evaluated on a plausible expected crowd state.")
    add_body(doc, "It does not by itself establish stampede prevention. The ATC data represent ordinary shopping-center movement, not labeled emergency crowd crush events. A safety claim should therefore be framed as a decision-support and counterfactual simulation capability, validated for normal-flow forecasting. Emergency validation requires scenario-based stress tests, expert-defined thresholds, and, where possible, additional event or evacuation data.")

    add_heading(doc, "Recommended next validation layer")
    add_bullets(doc, [
        "Run controlled surge, counterflow, and blocked-exit scenarios from the same observed seeds.",
        "Report safety indicators alongside trajectory error: density hotspots, local speed variance, counterflow conflict rate, exit throughput, and time above a chosen density threshold.",
        "Perform sensitivity analysis over agent count, desired speed, route closures, and interaction strength.",
        "Define operational alert thresholds with a domain expert before claiming automated intervention or prevention.",
        "Repeat the held-out evaluation across several test days and publish confidence intervals, not a single-day score alone.",
    ])

    add_heading(doc, "Reproducibility")
    add_body(doc, "Source dataset: ATC pedestrian tracking CSV files, using 10 available daily files in the local project. Development day: atc-20121111.csv. Held-out test day: atc-20121114.csv. Forecast snapshots: 6 per horizon, chosen across the usable span of the held-out day. History window: 5 seconds. Simulation seed: 42. The report values are generated by validation_runner.py and calibration is generated by calibrate_social_force.py.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("ATC Crowd Safety Explorer  |  Normal Flow Validation")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(90, 90, 90)

    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
